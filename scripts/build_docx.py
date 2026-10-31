#!/usr/bin/env python3
"""Build the GRHS review DOCX from canonical Markdown."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "SPECIFICATION.md"
OUTPUT = ROOT / "GRHS-v0.4.0.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(45, 45, 45)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(0, 70, 127)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            if col_index < len(row):
                add_inline(paragraph, row[col_index])
            if row_index == 0:
                # Do not strand a repeated table header at the foot of a page;
                # keep it with the first data row whenever Word paginates.
                paragraph.paragraph_format.keep_with_next = True
                set_cell_shading(cell, "D9D9D9")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc(document: Document) -> None:
    heading = document.add_paragraph(style="Title Small")
    heading.paragraph_format.space_before = Pt(14)
    heading.add_run("Contents")
    set_keep_with_next(heading)
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update this field if the contents are not visible."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    document.add_page_break()


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Title", 28, 0, 10),
        ("Subtitle", 15, 0, 14),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Subtitle" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, level, size, before, after in (
        ("GRHS Heading 1", 0, 17, 14, 6),
        ("GRHS Heading 2", 1, 13, 11, 4),
        ("GRHS Heading 3", 2, 11, 8, 3),
    ):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH) if name not in styles else styles[name]
        style.base_style = styles["Normal"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        p_pr = style.element.get_or_add_pPr()
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            p_pr.append(outline)
        outline.set(qn("w:val"), str(level))

    if "Title Small" not in styles:
        small = styles.add_style("Title Small", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = styles["Title Small"]
    small.font.name = "Aptos Display"
    small.font.size = Pt(20)
    small.font.bold = True
    small.font.color.rgb = RGBColor(0, 0, 0)
    small.paragraph_format.space_after = Pt(8)
    small.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(2)

    # Keep the complete three-level contents list on one page. Word's default
    # TOC styles are unnecessarily loose for a compact technical standard.
    for toc_name, size, left in (
        ("TOC 1", 9.0, 0.0),
        ("TOC 2", 8.7, 0.18),
        ("TOC 3", 8.4, 0.36),
    ):
        if toc_name not in styles:
            continue
        style = styles[toc_name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.left_indent = Inches(left)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(70)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("GRHS")
    run.font.name = "Aptos Display"
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("A Modern Phonemic Latin Orthography for Hindustani")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)

    for line, bold in (
        ("Version 0.4.0", True),
        ("Public review draft", True),
        ("31 October 2026", False),
        ("Gary Bajaj", False),
    ):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(11)

    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(52)
    notice.paragraph_format.left_indent = Inches(0.7)
    notice.paragraph_format.right_indent = Inches(0.7)
    notice.add_run(
        "Individual proposal for technical review. Publication does not imply endorsement, "
        "adoption, registration, or approval by any external body."
    ).italic = True

    license_p = document.add_paragraph()
    license_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    license_p.paragraph_format.space_before = Pt(18)
    license_p.add_run("Copyright © 2026 Gary Bajaj • AGPL-3.0-only").font.size = Pt(9)
    document.add_page_break()


def add_document_control(document: Document) -> None:
    heading = document.add_paragraph(style="Title Small")
    heading.add_run("Document control")
    rows = [
        ["Field", "Value"],
        ["Version", "0.4.0"],
        ["Date", "31 October 2026"],
        ["Status", "Public review draft; not a final standard"],
        ["Author", "Gary Bajaj"],
        ["Reference variety", "Educated Delhi/Doabi Hindustani"],
        ["Canonical normalization", "Unicode NFC"],
        ["License", "AGPL-3.0-only"],
    ]
    add_table(document, rows)
    p = document.add_paragraph()
    p.add_run("Canonical source: ").bold = True
    p.add_run("SPECIFICATION.md in the tagged release repository.")


def add_markdown_body(document: Document, text: str) -> None:
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. Scope"))
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.15)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_together = True
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F2F2F2")
                p_pr.append(shd)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Cascadia Mono"
                r.font.size = Pt(8.7)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(document, rows)
            continue
        if line.startswith("#### "):
            p = document.add_paragraph(style="GRHS Heading 3")
            add_inline(p, line[5:])
            i += 1
            continue
        if line.startswith("### "):
            p = document.add_paragraph(style="GRHS Heading 2")
            add_inline(p, line[4:])
            i += 1
            continue
        if line.startswith("## "):
            p = document.add_paragraph(style="GRHS Heading 1")
            add_inline(p, line[3:])
            i += 1
            continue
        if line.startswith("- "):
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            # Use literal authored numbers in ordinary paragraphs. This avoids
            # Word continuing automatic counters between unrelated sections,
            # while retaining one properly spaced requirement per paragraph.
            while i < len(lines):
                item = re.match(r"^(\d+)\.\s+(.*)$", lines[i])
                if item is None:
                    break
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
                p.paragraph_format.space_after = Pt(1)
                p.add_run(f"{item.group(1)}. ")
                add_inline(p, item.group(2))
                i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip() or nxt.startswith(("## ", "### ", "#### ", "```", "- ", "|")):
                break
            if re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = document.add_paragraph()
        add_inline(p, " ".join(part.strip() for part in paragraph_lines))


def build() -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    properties = document.core_properties
    properties.title = "GRHS 0.4.0 — A Modern Phonemic Latin Orthography for Hindustani"
    properties.subject = "Public review draft"
    properties.author = "Gary Bajaj"
    properties.last_modified_by = "Gary Bajaj"
    properties.created = datetime(2026, 10, 31, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 10, 31, tzinfo=timezone.utc)
    properties.keywords = "Hindustani, Hindi, Urdu, orthography, Latin, GRHS"
    properties.comments = "Generated from canonical SPECIFICATION.md"

    header = section.header.paragraphs[0]
    header.text = "GRHS 0.4.0  •  Public review draft"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(80, 80, 80)
    add_page_number(section.footer.paragraphs[0])

    add_cover(document)
    add_document_control(document)
    add_toc(document)
    add_markdown_body(document, SOURCE.read_text(encoding="utf-8"))

    document.save(OUTPUT)
    print(f"WROTE {OUTPUT}")


if __name__ == "__main__":
    build()

